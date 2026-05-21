#!/usr/bin/env python3
"""
SecureKernel Project Report Generator
Produces: SecureKernel_Project_Report.docx (70+ pages)
"""

import sys
import os

sys.path.insert(0, os.path.dirname(__file__))

from docx import Document
from docx.shared import Pt, Cm, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

from gen_part1 import add_frontmatter_and_ch1_to_4
from gen_part2 import add_ch5_to_7
from gen_part3 import add_ch8_to_refs


def setup_document():
    doc = Document()

    # --- Page setup: A4, standard academic margins ---
    for section in doc.sections:
        section.page_height = Cm(29.7)
        section.page_width  = Cm(21.0)
        section.top_margin    = Cm(2.54)
        section.bottom_margin = Cm(2.54)
        section.left_margin   = Cm(3.17)
        section.right_margin  = Cm(2.54)

    # --- Default Normal style: Times New Roman 12pt ---
    style = doc.styles['Normal']
    font  = style.font
    font.name = 'Times New Roman'
    font.size = Pt(12)
    pf = style.paragraph_format
    pf.space_before = Pt(0)
    pf.space_after  = Pt(6)
    pf.line_spacing = Pt(18)

    # --- Heading styles ---
    for h_name, h_size in [('Heading 1', 14), ('Heading 2', 12), ('Heading 3', 12)]:
        if h_name in doc.styles:
            s = doc.styles[h_name]
            s.font.name = 'Times New Roman'
            s.font.size = Pt(h_size)
            s.font.bold = True
            s.font.color.rgb = None   # black
            s.paragraph_format.space_before = Pt(12)
            s.paragraph_format.space_after  = Pt(6)

    return doc


def add_page_numbers(doc):
    """Add page numbers to the footer of all sections."""
    for section in doc.sections:
        footer = section.footer
        footer.is_linked_to_previous = False
        p = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.clear()

        run = p.add_run()
        run.font.name = 'Times New Roman'
        run.font.size = Pt(10)

        # Insert PAGE field
        fldChar1 = OxmlElement('w:fldChar')
        fldChar1.set(qn('w:fldCharType'), 'begin')
        instrText = OxmlElement('w:instrText')
        instrText.text = ' PAGE '
        fldChar2 = OxmlElement('w:fldChar')
        fldChar2.set(qn('w:fldCharType'), 'end')

        run._r.append(fldChar1)
        run._r.append(instrText)
        run._r.append(fldChar2)


def main():
    output_path = os.path.join(
        os.path.dirname(__file__),
        'SecureKernel_Project_Report.docx'
    )

    print("Creating document...")
    doc = setup_document()

    print("Adding front matter and Chapters 1-4...")
    add_frontmatter_and_ch1_to_4(doc)

    print("Adding Chapters 5-7...")
    add_ch5_to_7(doc)

    print("Adding Chapters 8-10 and References...")
    add_ch8_to_refs(doc)

    print("Adding page numbers...")
    add_page_numbers(doc)

    print(f"Saving to {output_path} ...")
    doc.save(output_path)

    size_mb = os.path.getsize(output_path) / (1024 * 1024)
    print(f"Done! File saved: {output_path} ({size_mb:.2f} MB)")
    print("Open in LibreOffice Writer or Microsoft Word to view.")


if __name__ == '__main__':
    main()
